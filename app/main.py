import io, logging
from datetime import timedelta, date
import fitz
from docx import Document as DocxDocument
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from .config import get_settings
from .database import Base, engine, get_db
from .models.core import *
from .schemas import *
from .security import hash_password,verify_password,create_token,decode_token
from .dependencies import current_user,admin_user
from .services import process_document,search,update_learning_state,analyze_mistake,readiness

logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s'); log=logging.getLogger('vectra')
app=FastAPI(title='VECTRA API',version='1.0.0',description='Backend for the proposed Vector-Enhanced Cognitive Tutoring & Reasoning Architecture research framework.')
app.add_middleware(CORSMiddleware,allow_origins=get_settings().cors_origin_list,allow_credentials=True,allow_methods=['*'],allow_headers=['*'])
@app.exception_handler(HTTPException)
async def http_error(_,exc):return JSONResponse(status_code=exc.status_code,content={'error':{'code':str(exc.status_code),'message':str(exc.detail),'details':{}}})
@app.exception_handler(Exception)
async def unexpected(_,exc): log.exception('unhandled error');return JSONResponse(status_code=500,content={'error':{'code':'INTERNAL_ERROR','message':'An unexpected server error occurred.','details':{}}})
@app.on_event('startup')
def startup(): Base.metadata.create_all(engine)
def token_pair(user):
 s=get_settings();return TokenOut(access_token=create_token(str(user.id),'access',timedelta(minutes=s.access_token_expire_minutes)),refresh_token=create_token(str(user.id),'refresh',timedelta(days=s.refresh_token_expire_days)))

@app.post('/api/auth/register',response_model=TokenOut,tags=['Authentication'],status_code=201)
def register(data:RegisterIn,db:Session=Depends(get_db)):
 if db.scalar(select(User).where(User.email==data.email.lower())):raise HTTPException(409,'Email is already registered')
 user=User(name=data.name,email=data.email.lower(),password_hash=hash_password(data.password));db.add(user);db.flush();db.add(StudentProfile(user_id=user.id));db.commit();return token_pair(user)
@app.post('/api/auth/login',response_model=TokenOut,tags=['Authentication'])
def login(data:LoginIn,db:Session=Depends(get_db)):
 user=db.scalar(select(User).where(User.email==data.email.lower()))
 if not user or not verify_password(data.password,user.password_hash):raise HTTPException(401,'Invalid email or password')
 return token_pair(user)
@app.get('/api/auth/me',response_model=UserOut,tags=['Authentication'])
def me(user:User=Depends(current_user)):return user
@app.post('/api/auth/refresh',response_model=TokenOut,tags=['Authentication'])
def refresh(refresh_token:str,db:Session=Depends(get_db)):
 try: user=db.get(User,decode_token(refresh_token,'refresh'))
 except Exception:raise HTTPException(401,'Invalid refresh token')
 if not user:raise HTTPException(401,'User not found')
 return token_pair(user)

@app.post('/api/documents/upload',tags=['Documents'],status_code=201)
def upload(file:UploadFile=File(...),user:User=Depends(admin_user),db:Session=Depends(get_db)):
 ext=(file.filename or '').rsplit('.',1)[-1].lower();allowed={'pdf','docx','txt'}
 if ext not in allowed:raise HTTPException(415,'Only PDF, DOCX and TXT are supported')
 data=file.file.read()
 if len(data)>get_settings().max_upload_size_mb*1024*1024:raise HTTPException(413,'File exceeds configured size limit')
 try:
  text=fitz.open(stream=data,filetype='pdf').get_text() if ext=='pdf' else '\n'.join(p.text for p in DocxDocument(io.BytesIO(data)).paragraphs) if ext=='docx' else data.decode('utf-8')
 except Exception:raise HTTPException(422,'Document text extraction failed')
 doc=Document(uploaded_by=user.id,filename=file.filename,file_type=ext);db.add(doc);db.flush();process_document(db,doc,text);return {'id':doc.id,'status':doc.status,'filename':doc.filename}
@app.get('/api/documents',tags=['Documents'])
def documents(user:User=Depends(current_user),db:Session=Depends(get_db)):return db.scalars(select(Document).where(Document.uploaded_by==user.id) if user.role==Role.STUDENT else select(Document)).all()
@app.get('/api/documents/{document_id}',tags=['Documents'])
def document(document_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
 d=db.get(Document,document_id)
 if not d or (user.role==Role.STUDENT and d.uploaded_by!=user.id):raise HTTPException(404,'Document not found')
 return d
@app.delete('/api/documents/{document_id}',status_code=204,tags=['Documents'])
def delete_document(document_id:int,user:User=Depends(admin_user),db:Session=Depends(get_db)):
 d=db.get(Document,document_id)
 if not d:raise HTTPException(404,'Document not found')
 db.delete(d);db.commit()
@app.post('/api/search',tags=['Search'])
def vector_search(data:SearchIn,user:User=Depends(current_user),db:Session=Depends(get_db)):return {'results':search(db,data.query,data.top_k,data.topic)}

@app.post('/api/concepts',tags=['Concepts'],status_code=201)
def create_concept(data:ConceptIn,user:User=Depends(admin_user),db:Session=Depends(get_db)):
 if db.scalar(select(Concept).where(Concept.name==data.name)):raise HTTPException(409,'Concept already exists')
 c=Concept(**data.model_dump());db.add(c);db.commit();db.refresh(c);return c
@app.get('/api/concepts',tags=['Concepts'])
def concepts(user:User=Depends(current_user),db:Session=Depends(get_db)):return db.scalars(select(Concept)).all()
@app.get('/api/concepts/{concept_id}',tags=['Concepts'])
def concept(concept_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
 c=db.get(Concept,concept_id)
 if not c:raise HTTPException(404,'Concept not found')
 return c
@app.post('/api/concepts/{concept_id}/relationships',tags=['Concepts'],status_code=201)
def relate(concept_id:int,data:RelationshipIn,user:User=Depends(admin_user),db:Session=Depends(get_db)):
 if not db.get(Concept,concept_id) or not db.get(Concept,data.target_concept_id):raise HTTPException(404,'Concept not found')
 r=ConceptRelationship(source_concept_id=concept_id,**data.model_dump());db.add(r);db.commit();return r
def relations(concept_id,kind,db):
 rows=db.scalars(select(ConceptRelationship).where(ConceptRelationship.source_concept_id==concept_id,ConceptRelationship.relationship_type==kind)).all();return [{'relationship':r.relationship_type,'weight':r.weight,'concept':db.get(Concept,r.target_concept_id)} for r in rows]
@app.get('/api/concepts/{concept_id}/related',tags=['Concepts'])
def related(concept_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):return relations(concept_id,RelationshipType.RELATED,db)
@app.get('/api/concepts/{concept_id}/prerequisites',tags=['Concepts'])
def prerequisites(concept_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):return relations(concept_id,RelationshipType.PREREQUISITE,db)

@app.post('/api/tutor/chat',tags=['Tutor'])
def chat(data:TutorIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 sources=search(db,data.message,5,data.topic); concepts=[s['topic'] for s in sources if s['topic']]
 if not sources:return {'answer':'I do not have retrieved course material for this yet. Try uploading approved learning resources or ask for guided practice.','sources':[],'concepts':[],'recommended_action':'Upload or select learning material.','confidence':0.0}
 context='\n\n'.join(s['chunk'][:700] for s in sources); state=''
 if concepts:
  c=db.scalar(select(Concept).where(Concept.name==concepts[0])); ls=c and db.scalar(select(LearningState).where(LearningState.user_id==user.id,LearningState.concept_id==c.id)); state=f' Your current mastery is {ls.mastery_score:.0f}/100.' if ls else ''
 return {'answer':f'Based on the retrieved learning material: {context[:1200]}{state} Work through the key idea, then attempt a similar question before requesting a full solution.','sources':sources,'concepts':list(dict.fromkeys(concepts)),'recommended_action':'Attempt one guided practice question.','confidence':round(max(s['similarity'] for s in sources),2)}
@app.post('/api/tutor/hint',tags=['Tutor'])
def hint(data:HintIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 q=db.get(Question,data.question_id)
 if not q:raise HTTPException(404,'Question not found')
 levels=['Try identifying the quantities and what is being asked.','Focus on the relationship between the given values.','Use the core formula for this concept.','Apply the formula one operation at a time.',q.explanation,'Now solve a closely related variant.']
 return {'level':data.level,'hint':levels[data.level],'reveals_solution':data.level>=4}
@app.post('/api/tutor/explain',tags=['Tutor'])
def explain(data:HintIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 q=db.get(Question,data.question_id)
 if not q:raise HTTPException(404,'Question not found')
 return {'explanation':q.explanation if data.level>=4 else 'Use progressive hints before revealing the complete solution.','level':data.level}
@app.post('/api/tutor/similar-question',tags=['Tutor'])
def similar(data:HintIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 q=db.get(Question,data.question_id)
 if not q:raise HTTPException(404,'Question not found')
 other=db.scalar(select(Question).where(Question.concept_id==q.concept_id,Question.id!=q.id))
 if not other:raise HTTPException(404,'No validated similar question is available')
 return other

@app.post('/api/questions',tags=['Admin'],status_code=201)
def question_create(data:QuestionIn,user:User=Depends(admin_user),db:Session=Depends(get_db)):
 if data.correct_answer not in data.options:raise HTTPException(422,'Correct answer must be one of the options')
 q=Question(**data.model_dump());db.add(q);db.commit();db.refresh(q);return q
@app.post('/api/practice/start',tags=['Practice'])
def practice_start(data:PracticeStartIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 qs=db.scalars(select(Question).where(Question.approved==True)).all()
 if data.topic: qs=[q for q in qs if (c:=db.get(Concept,q.concept_id)) and c.name==data.topic]
 return {'session_id':f'{user.id}-{int(__import__("time").time())}','next_question':qs[0] if qs else None,'initial_difficulty':'MEDIUM'}
@app.post('/api/practice/{session_id}/next',tags=['Practice'])
def practice_next(session_id:str,user:User=Depends(current_user),db:Session=Depends(get_db)):
 q=db.scalar(select(Question).where(Question.approved==True));return {'session_id':session_id,'question':q}
@app.post('/api/practice/{session_id}/answer',tags=['Practice'])
def practice_answer(session_id:str,data:AnswerIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 q=db.get(Question,data.question_id)
 if not q:raise HTTPException(404,'Question not found')
 attempt=QuestionAttempt(user_id=user.id,**data.model_dump(),is_correct=data.selected_answer.strip()==q.correct_answer.strip());db.add(attempt);db.flush(); mistake=analyze_mistake(db,attempt,q); state=update_learning_state(db,attempt,q);db.commit()
 difficulty='HARD' if attempt.is_correct and data.solving_time_seconds<90 else 'EASY' if not attempt.is_correct else 'MEDIUM'
 return {'attempt_id':attempt.id,'is_correct':attempt.is_correct,'explanation':q.explanation,'mistake':mistake,'learning_state':state,'recommended_difficulty':difficulty}
@app.get('/api/practice/{session_id}/summary',tags=['Practice'])
def practice_summary(session_id:str,user:User=Depends(current_user),db:Session=Depends(get_db)):
 a=db.scalars(select(QuestionAttempt).where(QuestionAttempt.user_id==user.id)).all();return {'session_id':session_id,'attempts':len(a),'accuracy':sum(x.is_correct for x in a)/len(a) if a else 0}

@app.post('/api/revision/generate',tags=['Revision'])
def revision_generate(data:RevisionGenerateIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 states=db.scalars(select(LearningState).where(LearningState.user_id==user.id).order_by(LearningState.mastery_score)).all(); plans=[]
 for i,s in enumerate(states[:5],1):
  p=RevisionPlan(user_id=user.id,concept_id=s.concept_id,priority=i,scheduled_date=data.scheduled_date or date.today(),recommended_activity='Review formulas, then complete a timed mini-practice set.');db.add(p);plans.append(p)
 db.commit();return plans
@app.get('/api/revision/plan',tags=['Revision'])
def revision_plan(user:User=Depends(current_user),db:Session=Depends(get_db)):return db.scalars(select(RevisionPlan).where(RevisionPlan.user_id==user.id).order_by(RevisionPlan.priority)).all()
@app.post('/api/revision/{plan_id}/complete',tags=['Revision'])
def revision_complete(plan_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
 p=db.get(RevisionPlan,plan_id)
 if not p or p.user_id!=user.id:raise HTTPException(404,'Revision plan not found')
 p.completion_status=True;db.commit();return p
@app.post('/api/readiness/recalculate',tags=['Readiness'])
def readiness_recalculate(company:str,user:User=Depends(current_user),db:Session=Depends(get_db)):return readiness(db,user.id,company)
@app.get('/api/readiness',tags=['Readiness'])
def readiness_all(user:User=Depends(current_user),db:Session=Depends(get_db)):return db.scalars(select(Readiness).where(Readiness.user_id==user.id)).all()
@app.get('/api/readiness/{company}',tags=['Readiness'])
def readiness_one(company:str,user:User=Depends(current_user),db:Session=Depends(get_db)):return readiness(db,user.id,company)

@app.get('/api/analytics/overview',tags=['Analytics'])
def overview(user:User=Depends(current_user),db:Session=Depends(get_db)):
 a=db.scalars(select(QuestionAttempt).where(QuestionAttempt.user_id==user.id)).all(); states=db.scalars(select(LearningState).where(LearningState.user_id==user.id)).all(); mistakes=db.scalars(select(Mistake).where(Mistake.user_id==user.id)).all()
 return {'overall_accuracy':round(100*sum(x.is_correct for x in a)/len(a),2) if a else 0,'total_questions':len(a),'average_solving_time':round(sum(x.solving_time_seconds for x in a)/len(a),2) if a else 0,'mastered_concepts':[s.concept_id for s in states if s.mastery_score>=70],'weak_concepts':[s.concept_id for s in states if s.mastery_score<50],'mistake_distribution':{x.value:sum(m.mistake_type==x for m in mistakes) for x in MistakeType},'readiness_scores':db.scalars(select(Readiness).where(Readiness.user_id==user.id)).all()}
@app.get('/api/analytics/topics',tags=['Analytics'])
def topics(user:User=Depends(current_user),db:Session=Depends(get_db)):
 return [{'concept':db.get(Concept,s.concept_id).name,'mastery':s.mastery_score,'accuracy':s.accuracy,'average_time':s.average_time,'trend':s.trend} for s in db.scalars(select(LearningState).where(LearningState.user_id==user.id)).all()]
@app.get('/api/analytics/mistakes',tags=['Analytics'])
def analytics_mistakes(user:User=Depends(current_user),db:Session=Depends(get_db)):
 m=db.scalars(select(Mistake).where(Mistake.user_id==user.id)).all();return {x.value:sum(i.mistake_type==x for i in m) for x in MistakeType}
@app.post('/api/mock-tests',tags=['Admin'],status_code=201)
def mock_create(data:MockTestIn,user:User=Depends(admin_user),db:Session=Depends(get_db)):
 m=MockTest(title=data.title,company=data.company,difficulty=data.difficulty,duration_minutes=data.duration_minutes,question_count=data.question_count,config=data.config);db.add(m);db.flush()
 for i,qid in enumerate(data.question_ids):db.add(MockTestQuestion(mock_test_id=m.id,question_id=qid,sequence=i+1))
 db.commit();return m
@app.post('/api/mock-tests/{mock_test_id}/start',tags=['Mock Tests'])
def mock_start(mock_test_id:int,user:User=Depends(current_user),db:Session=Depends(get_db)):
 m=db.get(MockTest,mock_test_id)
 if not m:raise HTTPException(404,'Mock test not found')
 ids=db.scalars(select(MockTestQuestion.question_id).where(MockTestQuestion.mock_test_id==m.id).order_by(MockTestQuestion.sequence)).all();return {'mock_test':m,'questions':[db.get(Question,x) for x in ids]}
@app.post('/api/mock-tests/{mock_test_id}/submit',tags=['Mock Tests'])
def mock_submit(mock_test_id:int,data:MockSubmitIn,user:User=Depends(current_user),db:Session=Depends(get_db)):
 m=db.get(MockTest,mock_test_id)
 if not m:raise HTTPException(404,'Mock test not found')
 correct=0;times=[]
 for answer in data.answers:
  q=db.get(Question,answer.question_id)
  if not q:continue
  ok=answer.selected_answer==q.correct_answer;correct+=ok;times.append(answer.solving_time_seconds);a=QuestionAttempt(user_id=user.id,**answer.model_dump(),is_correct=ok);db.add(a);db.flush();analyze_mistake(db,a,q);update_learning_state(db,a,q)
 accuracy=100*correct/len(data.answers) if data.answers else 0;attempt=MockTestAttempt(user_id=user.id,mock_test_id=m.id,score=correct,accuracy=accuracy,average_time=sum(times)/len(times) if times else 0);db.add(attempt);db.commit();r=readiness(db,user.id,m.company);return {'score':correct,'accuracy':accuracy,'average_time':attempt.average_time,'readiness':r}
